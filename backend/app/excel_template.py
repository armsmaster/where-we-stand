"""Генерация файла-шаблона (ТЗ 4.7).

Шаблон собирается кодом, а не хранится статическим файлом, чтобы
не рассинхронизироваться со спецификацией столбцов.
"""

from __future__ import annotations

import datetime as dt
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .data_model import MAX_ROWS

HEADERS = ["key", "label", "unit", "min", "max", "current", "prev", "decimals", "inverted"]

EXAMPLE_ROWS = [
    ["imoex", "Индекс МосБиржи", "п.", 2512, 3521, 2890, 2845, 0, False],
    ["ofz_10y", "Доходность ОФЗ 10 лет", "%", 13.8, 17.2, 15.1, 15.4, 1, True],
    ["corp_spread", "Кредитный спред", "б.п.", 180, 520, 340, 355, 0, True],
    ["brent", "Нефть Brent", "$/барр.", 62, 88, 71.4, 69.8, 1, False],
    ["gold", "Золото", "$/унц.", 1980, 2790, 2620, 2585, 0, False],
]

COLUMN_DOCS = [
    ("key", "да", "строка", "Латиница, цифры, подчёркивание. 1–32 символа. Уникален в файле. "
                            "По нему настройки привязывают индивидуальное оформление строки"),
    ("label", "да", "строка", "1–40 символов. Подпись строки в картинке"),
    ("unit", "нет", "строка", "0–16 символов. Единица измерения, выводится рядом с подписью"),
    ("min", "да", "число", "Годовой минимум"),
    ("max", "да", "число", "Годовой максимум"),
    ("current", "да", "число", "Текущий уровень. Должен лежать в диапазоне min…max"),
    ("prev", "да", "число", "Уровень неделю назад. Должен лежать в диапазоне min…max"),
    ("decimals", "нет", "целое", "0–3. Знаков после запятой при выводе значений. По умолчанию 0"),
    ("inverted", "нет", "логический", "ИСТИНА — показатель инвертирован по смыслу (доходность, "
                                      "премия). Влияет только на сноску, не на геометрию"),
]

RULES = [
    "Лист `data` обязателен, имя пишется строчными латинскими буквами.",
    "Лист `meta` обязателен: в A1 ключ `as_of`, в B1 — дата актуальности данных.",
    f"Строк данных от 1 до {MAX_ROWS}. Штатная конфигурация — 5.",
    "Порядок строк в картинке повторяет порядок строк в файле. Сервис не сортирует данные.",
    "Пустая строка прерывает чтение: всё, что ниже неё, в картинку не попадёт.",
    "Значения должны быть числами, а не текстом. Знак процента в ячейке недопустим.",
    "Если значения считаются формулами — сохраните файл в Excel перед загрузкой, "
    "иначе результат формулы не попадёт в файл.",
]

HEADER_FILL = PatternFill("solid", fgColor="F1EFE8")
FIXED_TIME = dt.datetime(2024, 1, 1)


def build() -> bytes:
    workbook = Workbook()

    data_sheet = workbook.active
    data_sheet.title = "data"
    data_sheet.append(HEADERS)
    for cell in data_sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for row in EXAMPLE_ROWS:
        data_sheet.append(row)

    widths = [14, 32, 10, 10, 10, 10, 10, 11, 11]
    for index, width in enumerate(widths, start=1):
        data_sheet.column_dimensions[get_column_letter(index)].width = width
    data_sheet.freeze_panes = "A2"

    meta_sheet = workbook.create_sheet("meta")
    meta_sheet["A1"] = "as_of"
    meta_sheet["B1"] = dt.date.today()
    meta_sheet["B1"].number_format = "DD.MM.YYYY"
    meta_sheet["A1"].font = Font(bold=True)
    meta_sheet["D1"] = "Дата актуальности данных. Попадёт в строку источника и в имя файла."
    meta_sheet.column_dimensions["A"].width = 14
    meta_sheet.column_dimensions["B"].width = 14
    meta_sheet.column_dimensions["D"].width = 70

    readme = workbook.create_sheet("readme")
    readme["A1"] = "Лист data — столбцы"
    readme["A1"].font = Font(bold=True, size=13)
    readme.append([])
    readme.append(["Столбец", "Обязательный", "Тип", "Правила"])
    for cell in readme[3]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for doc in COLUMN_DOCS:
        readme.append(list(doc))

    readme.append([])
    rules_row = readme.max_row + 1
    readme.cell(row=rules_row, column=1, value="Правила заполнения").font = Font(bold=True, size=13)
    for rule in RULES:
        readme.append([rule])

    readme.column_dimensions["A"].width = 16
    readme.column_dimensions["B"].width = 15
    readme.column_dimensions["C"].width = 14
    readme.column_dimensions["D"].width = 84
    for row in readme.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    workbook.properties.created = FIXED_TIME
    workbook.properties.modified = FIXED_TIME
    workbook.properties.creator = "Где мы находимся"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
