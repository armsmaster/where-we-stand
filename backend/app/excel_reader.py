"""Чтение и валидация файла данных (ТЗ 4, 5).

Весь файл обрабатывается целиком, все проблемы собираются разом:
останавливаться на первой ошибке запрещено (ТЗ 5.3).
"""

from __future__ import annotations

import datetime as dt
import difflib
import io
import re
from dataclasses import dataclass, field

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .data_model import EXPECTED_ROWS, KEY_PATTERN, MAX_ROWS, Meta, Row
from .issues import Issues
from .rules import check_range, check_row_warnings

DATA_SHEET = "data"
META_SHEET = "meta"

REQUIRED_COLUMNS = ("key", "label", "min", "max", "current", "prev")
OPTIONAL_COLUMNS = ("unit", "decimals", "inverted")
ALL_COLUMNS = REQUIRED_COLUMNS + OPTIONAL_COLUMNS
NUMERIC_COLUMNS = ("min", "max", "current", "prev")

ALLOWED_EXTENSIONS = (".xlsx", ".xlsm")

TRUE_WORDS = {"true", "истина", "да", "1", "yes", "y"}
FALSE_WORDS = {"false", "ложь", "нет", "0", "no", "n"}

# Пробельные символы, которые приходят при копировании из веба и ломают
# проверку ключа невнятным сообщением (ТЗ 4.3).
SPACE_CHARS = "   ​"  # NBSP, узкий NBSP, тонкий, нулевой ширины

KEY_RE = re.compile(KEY_PATTERN)


@dataclass
class ParseResult:
    issues: Issues
    rows: list[Row] = field(default_factory=list)
    meta: Meta | None = None

    @property
    def ok(self) -> bool:
        return self.issues.ok and self.meta is not None


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value)
    for char in SPACE_CHARS:
        text = text.replace(char, " ")
    return re.sub(r"\s+", " ", text).strip()


def looks_numeric(text: str) -> bool:
    candidate = text.replace(" ", "").replace(",", ".")
    try:
        float(candidate)
        return True
    except ValueError:
        return False


def parse(content: bytes, filename: str) -> ParseResult:
    issues = Issues()
    result = ParseResult(issues=issues)

    lowered = filename.lower()
    if not lowered.endswith(ALLOWED_EXTENSIONS):
        suffix = lowered.rsplit(".", 1)[-1] if "." in lowered else "без расширения"
        issues.error(
            "файл",
            f"Формат «.{suffix}» не поддерживается. Сохраните файл как .xlsx "
            "(в Excel: Файл → Сохранить как → Книга Excel).",
        )
        return result

    try:
        values_wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 — сообщение пользователю важнее типа
        issues.error(
            "файл",
            f"Не удалось открыть файл как книгу Excel: {exc}. "
            "Проверьте, что файл не повреждён и сохранён в формате .xlsx.",
        )
        return result

    try:
        sheet_names = list(values_wb.sheetnames)

        if DATA_SHEET not in sheet_names:
            listed = ", ".join(f"«{name}»" for name in sheet_names) or "нет ни одного"
            issues.error(
                f"лист {DATA_SHEET}",
                f"Не найден лист `data`. В файле есть листы: {listed}. "
                "Переименуйте нужный лист в `data` (регистр значим).",
            )
        else:
            formulas = _FormulaLookup(content)
            _read_data_sheet(values_wb[DATA_SHEET], issues, result, formulas)

        if META_SHEET not in sheet_names:
            issues.error(
                f"лист {META_SHEET}",
                "Не найден лист `meta` с ключом `as_of`. Добавьте лист `meta`, "
                "в ячейку A1 впишите `as_of`, в B1 — дату актуальности данных. "
                "Проще всего скачать шаблон и перенести данные в него.",
            )
        else:
            result.meta = _read_meta_sheet(values_wb[META_SHEET], issues)
    finally:
        values_wb.close()

    return result


class _FormulaLookup:
    """Второй проход по файлу: текст формул для сообщения о пустом кеше.

    openpyxl не вычисляет формулы — он читает результат, сохранённый Excel'ом.
    Если файл собран скриптом или ни разу не открывался в Excel, кеша нет
    и расчётные ячейки читаются как пустые. Без отдельного сообщения
    аналитик получит «ячейка пуста» на визуально заполненном файле (ТЗ 5.1).
    """

    def __init__(self, content: bytes) -> None:
        self._content = content
        self._cache: dict[str, str] | None = None

    def formula_at(self, coordinate: str) -> str | None:
        if self._cache is None:
            self._cache = {}
            try:
                wb = load_workbook(io.BytesIO(self._content), data_only=False, read_only=True)
            except Exception:  # noqa: BLE001 — подсказка необязательна
                return None
            try:
                if DATA_SHEET in wb.sheetnames:
                    sheet = wb[DATA_SHEET]
                    for row in sheet.iter_rows():
                        for cell in row:
                            if isinstance(cell.value, str) and cell.value.startswith("="):
                                letter = get_column_letter(cell.column)
                                self._cache[f"{letter}{cell.row}"] = cell.value
            finally:
                wb.close()
        return self._cache.get(coordinate)


def _read_data_sheet(sheet, issues: Issues, result: ParseResult, formulas: _FormulaLookup) -> None:
    grid = [list(row) for row in sheet.iter_rows(values_only=True)]
    if not grid:
        issues.error("лист data", "Лист `data` не содержит строк с данными.")
        return

    header_cells = grid[0]
    headers: dict[str, int] = {}
    seen: dict[str, int] = {}
    unknown: list[tuple[str, int]] = []

    for index, raw in enumerate(header_cells):
        name = normalize_text(raw)
        if not name:
            continue
        if name in ALL_COLUMNS:
            if name in seen:
                issues.error(
                    f"{get_column_letter(index + 1)}1",
                    f"Столбец `{name}` встречается дважды: "
                    f"{get_column_letter(seen[name] + 1)}1 и {get_column_letter(index + 1)}1. "
                    "Оставьте один.",
                )
                continue
            seen[name] = index
            headers[name] = index
        else:
            unknown.append((name, index))

    # ТЗ 5.1 помещает неизвестный столбец в таблицу блокирующих ошибок, но само
    # сообщение там заканчивается словами «будет проигнорирован». Блокировать отрисовку
    # из-за лишнего столбца с комментарием было бы абсурдом, поэтому — предупреждение.
    for name, index in unknown:
        address = f"{get_column_letter(index + 1)}1"
        issues.warning(
            address,
            f"Столбец `{name}` в ячейке {address} не входит в спецификацию "
            "и будет проигнорирован.",
        )

    for name in REQUIRED_COLUMNS:
        if name in headers:
            continue
        candidates = [
            (title, index)
            for title, index in unknown
            if difflib.SequenceMatcher(None, title.lower(), name).ratio() >= 0.6
        ]
        if candidates:
            title, index = candidates[0]
            address = f"{get_column_letter(index + 1)}1"
            issues.error(
                f"столбец {name}",
                f"Не найден обязательный столбец `{name}`. Возможно, имеется в виду "
                f"столбец `{title}` в ячейке {address} — переименуйте его в `{name}`.",
            )
        else:
            issues.error(
                f"столбец {name}",
                f"Не найден обязательный столбец `{name}`. Добавьте его в строку 1 "
                "листа `data` — имя пишется строго латиницей, в нижнем регистре.",
            )

    # О пропущенных столбцах уже сообщено выше, но разбор на этом не останавливается:
    # иначе один переименованный заголовок скрывает все ошибки в ячейках и исправление
    # файла превращается в многократный цикл загрузок (ТЗ 5.3).

    body_end = len(grid)
    for offset in range(1, len(grid)):
        if all(normalize_text(cell) == "" for cell in grid[offset]):
            body_end = offset
            break

    for offset in range(body_end, len(grid)):
        if any(normalize_text(cell) != "" for cell in grid[offset]):
            issues.warning(
                f"строка {body_end + 1}",
                f"Чтение остановлено на строке {body_end + 1} (пустая). "
                f"Ниже, начиная со строки {offset + 1}, есть непустые ячейки — "
                "они не попали в картинку.",
            )
            break

    body = grid[1:body_end]
    if not body:
        issues.error("лист data", "Лист `data` не содержит строк с данными.")
        return

    if len(body) > MAX_ROWS:
        issues.error(
            "лист data",
            f"В файле {len(body)} строк данных. Максимум — {MAX_ROWS}.",
        )
        return

    # Дубли ищем по сырым значениям, а не по успешно разобранным строкам:
    # иначе любая другая ошибка в строке скрывает дубликат ключа.
    if "key" in headers:
        keys_seen: dict[str, int] = {}
        for offset, raw_row in enumerate(body):
            excel_row = offset + 2
            key = normalize_text(_cell(raw_row, headers["key"]))
            if not key:
                continue
            if key in keys_seen:
                issues.error(
                    f"строка {excel_row}",
                    f"Значение `{key}` в столбце `key` встречается дважды: "
                    f"строки {keys_seen[key]} и {excel_row}. Ключи должны быть уникальны.",
                )
            else:
                keys_seen[key] = excel_row

    rows: list[Row] = []
    for offset, raw_row in enumerate(body):
        parsed = _read_row(raw_row, offset + 2, headers, issues, formulas)
        if parsed is not None:
            rows.append(parsed)

    result.rows = rows

    if len(body) != EXPECTED_ROWS:
        issues.warning(
            "лист data",
            f"В файле {len(body)} строк данных. Штатная конфигурация — "
            f"{EXPECTED_ROWS} классов активов.",
        )


def _cell(raw_row: list, index: int):
    return raw_row[index] if index < len(raw_row) else None


def _read_row(
    raw_row: list,
    excel_row: int,
    headers: dict[str, int],
    issues: Issues,
    formulas: _FormulaLookup,
) -> Row | None:
    errors_before = len(issues.errors)

    if "key" not in headers:
        key = ""
    else:
        key_index = headers["key"]
        key = _read_key(raw_row, key_index, excel_row, issues)

    label = _read_label(raw_row, headers, excel_row, key, issues) if "label" in headers else ""

    unit = ""
    if "unit" in headers:
        unit_index = headers["unit"]
        unit = normalize_text(_cell(raw_row, unit_index))
        if len(unit) > 16:
            address = f"{get_column_letter(unit_index + 1)}{excel_row}"
            issues.error(
                address,
                f"Ячейка {address}: единица измерения длиной {len(unit)} символов. "
                "Максимум — 16.",
            )
            unit = unit[:16]

    numbers: dict[str, float] = {}
    for name in NUMERIC_COLUMNS:
        if name not in headers:
            continue
        index = headers[name]
        address = f"{get_column_letter(index + 1)}{excel_row}"
        value = _read_number(
            _cell(raw_row, index), address, name, key or str(excel_row), issues, formulas
        )
        if value is not None:
            numbers[name] = value

    decimals = 0
    if "decimals" in headers:
        index = headers["decimals"]
        address = f"{get_column_letter(index + 1)}{excel_row}"
        raw = _cell(raw_row, index)
        if raw is not None and normalize_text(raw) != "":
            if isinstance(raw, bool) or not isinstance(raw, (int, float)) or raw != int(raw):
                issues.error(
                    address,
                    f"Ячейка {address}: ожидалось целое число от 0 до 3, "
                    f"найдено «{normalize_text(raw)}».",
                )
            elif not 0 <= int(raw) <= 3:
                issues.error(
                    address,
                    f"Ячейка {address}: ожидалось целое число от 0 до 3, найдено {int(raw)}.",
                )
            else:
                decimals = int(raw)

    inverted = False
    if "inverted" in headers:
        index = headers["inverted"]
        address = f"{get_column_letter(index + 1)}{excel_row}"
        raw = _cell(raw_row, index)
        if isinstance(raw, bool):
            inverted = raw
        elif raw is not None and normalize_text(raw) != "":
            word = normalize_text(raw).lower()
            if word in TRUE_WORDS:
                inverted = True
            elif word in FALSE_WORDS:
                inverted = False
            else:
                issues.error(
                    address,
                    f"Ячейка {address}: значение «{normalize_text(raw)}» не распознано. "
                    "Допустимы ИСТИНА, ЛОЖЬ, TRUE, FALSE, да, нет, 1, 0.",
                )

    row_label = f"Строка {excel_row} (`{key or excel_row}`)"
    check_range(
        numbers.get("min"), numbers.get("max"), numbers.get("current"), numbers.get("prev"),
        row_label, issues,
    )

    if len(issues.errors) != errors_before:
        return None
    if len(numbers) != len(NUMERIC_COLUMNS) or not key or not label:
        return None

    row = Row(
        key=key,
        label=label,
        unit=unit,
        min=numbers["min"],
        max=numbers["max"],
        current=numbers["current"],
        prev=numbers["prev"],
        decimals=decimals,
        inverted=inverted,
    )
    check_row_warnings(row, f"Строка {excel_row} (`{key}`)", issues)
    return row


def _read_key(raw_row: list, index: int, excel_row: int, issues: Issues) -> str:
    address = f"{get_column_letter(index + 1)}{excel_row}"
    key = normalize_text(_cell(raw_row, index))
    if not key:
        issues.error(address, f"Ячейка {address} пуста. Требуется значение `key`.")
    elif not KEY_RE.match(key):
        issues.error(
            address,
            f"Ячейка {address}: ключ «{key}» содержит недопустимые символы "
            "или слишком длинный. Допустимы латиница, цифры и подчёркивание, "
            "от 1 до 32 символов.",
        )
    return key


def _read_label(raw_row: list, headers: dict[str, int], excel_row: int, key: str,
                issues: Issues) -> str:
    index = headers["label"]
    address = f"{get_column_letter(index + 1)}{excel_row}"
    label = normalize_text(_cell(raw_row, index))
    if not label:
        issues.error(
            address,
            f"Ячейка {address} пуста. Для строки `{key or excel_row}` "
            "требуется значение `label`.",
        )
    elif len(label) > 40:
        issues.error(
            address,
            f"Ячейка {address}: подпись длиной {len(label)} символов. "
            "Максимум — 40.",
        )
    return label


def _read_number(
    raw: object,
    address: str,
    column: str,
    row_key: str,
    issues: Issues,
    formulas: _FormulaLookup,
) -> float | None:
    if isinstance(raw, bool):
        issues.error(
            address,
            f"Ячейка {address}: ожидалось число, найдено логическое значение. "
            f"Впишите числовое значение `{column}`.",
        )
        return None

    if isinstance(raw, (int, float)):
        return float(raw)

    if raw is None or normalize_text(raw) == "":
        formula = formulas.formula_at(address)
        if formula:
            issues.error(
                address,
                f"Ячейка {address} содержит формулу `{formula}`, но её результат "
                "в файле не сохранён. Откройте файл в Excel и сохраните ещё раз — "
                "либо вставьте в ячейку готовое значение.",
            )
        else:
            issues.error(
                address,
                f"Ячейка {address} пуста. Для строки `{row_key}` требуется "
                f"значение `{column}`.",
            )
        return None

    if isinstance(raw, (dt.datetime, dt.date)):
        issues.error(
            address,
            f"Ячейка {address}: ожидалось число, найдена дата. Смените формат "
            "ячейки на числовой (Главная → Числовой формат → Числовой).",
        )
        return None

    text = normalize_text(raw)
    if looks_numeric(text):
        issues.error(
            address,
            f"Ячейка {address}: значение выглядит числом, но сохранено как текст. "
            "В Excel: выделите столбец → Данные → Текст по столбцам → Готово.",
        )
    else:
        hint = (
            "Уберите знак процента и убедитесь, что в ячейке число, а не текст."
            if "%" in text
            else "Уберите лишние символы и убедитесь, что в ячейке число, а не текст."
        )
        issues.error(
            address,
            f"Ячейка {address}: ожидалось число, найден текст «{text}». {hint}",
        )
    return None


def _read_meta_sheet(sheet, issues: Issues) -> Meta | None:
    as_of: dt.date | None = None
    found_key = False

    for row_index, raw_row in enumerate(sheet.iter_rows(values_only=True), start=1):
        key = normalize_text(_cell(list(raw_row), 0)).lower()
        if not key:
            continue
        value = _cell(list(raw_row), 1)
        address = f"B{row_index}"

        if key == "as_of":
            found_key = True
            as_of = _parse_date(value, address, issues)
        else:
            issues.warning(
                f"A{row_index}",
                f"Лист `meta`: ключ `{key}` не используется и проигнорирован.",
            )

    if not found_key:
        issues.error(
            "лист meta",
            "На листе `meta` не найден ключ `as_of`. В ячейку A1 впишите `as_of`, "
            "в B1 — дату актуальности данных.",
        )
        return None

    return Meta(as_of=as_of) if as_of else None


def _parse_date(value: object, address: str, issues: Issues) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value

    text = normalize_text(value)
    if not text:
        issues.error(
            "лист meta",
            f"Лист `meta`, ячейка {address}: не указана дата актуальности данных. "
            "Впишите дату ячейкой формата «Дата» либо текстом вида 2026-03-12.",
        )
        return None

    try:
        return dt.date.fromisoformat(text)
    except ValueError:
        issues.error(
            "лист meta",
            f"Лист `meta`, ячейка {address}: значение «{text}» не распознано как дата. "
            "Впишите дату ячейкой формата «Дата» либо текстом вида 2026-03-12.",
        )
        return None
